"""Document ingestion: .txt / .pdf / .docx -> LangChain Documents -> chunks."""
from __future__ import annotations

from pathlib import Path

from langchain_community.document_loaders import PyPDFLoader, TextLoader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

SUPPORTED = {".txt", ".pdf", ".docx"}


def _load_docx(path: Path) -> list[Document]:
    import docx  # python-docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append("  ".join(c.text for c in row.cells))
    return [Document(page_content="\n".join(parts), metadata={"source": path.name})]


def load_documents(path: str | Path) -> list[Document]:
    path = Path(path)
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError(f"Unsupported file type {ext!r}; expected one of {sorted(SUPPORTED)}")
    if not path.exists():
        raise FileNotFoundError(path)
    if ext == ".txt":
        return TextLoader(str(path), encoding="utf-8").load()
    if ext == ".pdf":
        return PyPDFLoader(str(path)).load()
    return _load_docx(path)


def load_text(path: str | Path) -> str:
    """Full document text (pages joined by newlines)."""
    return "\n".join(d.page_content for d in load_documents(path))


def split_documents(docs: list[Document], chunk_size: int = 1000, chunk_overlap: int = 100) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(chunk_size=chunk_size, chunk_overlap=chunk_overlap)
    return splitter.split_documents(docs)


def load_and_split(path: str | Path, **kw) -> list[Document]:
    return split_documents(load_documents(path), **kw)
